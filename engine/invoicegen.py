"""Sun Orient invoice generator: K Line + ECL.

Claude reads the source documents and writes a small job JSON (numbers, PH shifts,
extra-work blocks). This module does the arithmetic and produces:
  - invoice PDF on letterhead
  - K Line IMS EDI CSV
  - invoice stack PDF (invoice + supporting docs)
"""
import csv, json, os, re, subprocess, datetime as dt
from decimal import Decimal, ROUND_HALF_UP
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader, PdfWriter

HERE = os.path.dirname(os.path.abspath(__file__))
RATES = json.load(open(os.path.join(HERE, "rates.json")))


def D(x):
    return Decimal(str(x))


def money(x):
    return D(x).quantize(Decimal("0.01"), ROUND_HALF_UP)


def fmt_money(x, dollar=True):
    s = f"{money(x):,.2f}"
    return f"${s}" if dollar else s


def fmt_qty(q):
    q = D(q)
    if q == q.to_integral():
        return f"{int(q):,}"
    return f"{q.normalize():f}" if abs(q.as_tuple().exponent) > 3 else f"{q:,.3f}"


# ---------------------------------------------------------------- K Line report parsing
def xls_to_xlsx(path, outdir):
    if path.lower().endswith(".xlsx"):
        return path
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", outdir, path],
                   check=True, capture_output=True)
    return os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".xlsx")


def parse_kline_report(path, workdir):
    import openpyxl
    ws = openpyxl.load_workbook(xls_to_xlsx(path, workdir), data_only=True)["Confirmation"]
    rows = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str):
            rows.setdefault(a.strip().upper(), r)
    v = lambda r, c: ws.cell(r, c).value

    def num(x):
        if x in (None, "", "-"):
            return D(0)
        if isinstance(x, str):
            m = re.search(r"[\d.]+", x.replace(",", ""))
            return D(m.group()) if m else D(0)
        return D(x)

    rep = {"vessel": v(9, 2), "voyage": str(v(9, 6)), "report_date": v(9, 9), "berth": v(11, 2),
           "berthed": v(11, 5), "est_unberth": v(11, 9), "cargo": {}}
    for code, label in [("MC", "MOTOR CARS"), ("CV", "COMMERCIAL VEHICLES"), ("HV", "HEAVY VEHICLES"), ("GC", "GENERAL CARGO (PKG)")]:
        r = rows[label]
        rep["cargo"][code] = {sec: (num(v(r, c)), num(v(r, c + 1))) for sec, c in [("DIS", 3), ("LOD", 5), ("SHF", 7), ("LNR", 9)]}
    r = rows["COMMENCED (HRS)"]
    rep["ops"] = {"premeet_start": v(r, 3), "cargo_start": v(r, 5), "cargo_end": v(r + 1, 5)}
    r = rows["LASHING / UNLASHING GANG"]
    rep["gang"] = {f"{s}D{d}": num(v(r, 3 + (s - 1) * 2 + (d - 1))) for s in (1, 2, 3) for d in (1, 2)}
    r = rows["SUPERVISOR / CHECKER"]
    rep["sup"] = {f"{s}D{d}": num(v(r, 3 + (s - 1) * 2 + (d - 1))) for s in (1, 2, 3) for d in (1, 2)}
    # first (UNIT / AMOUNT) row = equipment, second = towing/locksmith/fuel/others
    unit_rows = sorted(rr for rr in range(1, ws.max_row + 1) if str(v(rr, 1) or "").strip().upper() == "(UNIT / AMOUNT)")
    r1, r2 = unit_rows[0], unit_rows[1]
    rep["equip"] = {"MECH": num(v(r1, 3)), "FL4_OP": num(v(r1, 4)), "FL5": num(v(r1, 5)), "FL7": num(v(r1, 6)),
                    "FL16": num(v(r1, 7)), "FL_OP": num(v(r1, 8)), "TM_OP": num(v(r1, 9)), "TRAILER": num(v(r1, 10))}
    rep["misc"] = {"TOWING": (num(v(r2, 3)), num(v(r2, 4))), "LOCK": (num(v(r2, 5)), num(v(r2, 6))),
                   "FUEL": (num(v(r2, 7)), num(v(r2, 8))), "OTHERS": (num(v(r2, 9)), num(v(r2, 10)))}
    r = rows["MANHOUR / UNIT"]
    rep["nt_tugmaster"] = num(v(r, 10))
    r = rows["REMARK:"]
    rep["remarks"] = [str(v(rr, 2)).strip() for rr in range(r, r + 8) if v(rr, 2) not in (None, "")]
    return rep


# ---------------------------------------------------------------- K Line line items
SEC_WORD = {"DIS": "Discharged", "LOD": "Loaded", "SHF": "Shifted", "LNR": "Land/Reshipped"}


def line(key, desc, qty, unit_lbl, rate, tariff, amount=None, edi_qty=None, edi_rate=None, edi_unit=None):
    amt = money(amount if amount is not None else D(rate) * D(qty))
    return {"key": key, "desc": desc, "qty": D(qty), "unit": unit_lbl, "rate": D(rate), "amount": amt,
            "tariff": tariff, "edi_qty": D(edi_qty if edi_qty is not None else qty),
            "edi_rate": D(edi_rate if edi_rate is not None else rate), "edi_unit": edi_unit}


def svc_line(key, qty, ph=False, rate=None, desc=None):
    s = RATES["kline"]["services"][key]
    r = D(rate if rate is not None else s["rate"]) * (2 if ph else 1)
    d = (desc or s["desc"]) + (" (Sun/PH)" if ph else "")
    return line(key, d, qty, "Unit" if D(qty) == 1 else "Units", r, s["tariff"], edi_unit=s["unit"])


def kline_lines(rep, job):
    K = RATES["kline"]
    ph = set(job.get("ph_shifts", []))       # e.g. {"2D1","3D1"}
    ded = {k: D(v) for k, v in job.get("deduct", {}).items()}   # qty moved into extra blocks
    main = []
    for sec in ("DIS", "LOD", "SHF", "LNR"):
        for code in ("MC", "CV", "HV", "GC"):
            units, m3 = rep["cargo"][code][sec]
            if m3 == 0:
                continue
            t = K["cargo_tariff"][sec][code]
            main.append(line(f"{code}_{sec}", f"{fmt_qty(units)} Units of {K['cargo_names'][code]} ({SEC_WORD[sec]})",
                             m3, "M3", K["cargo_rates"][code], t, edi_unit="M3"))
    for pre, src in (("LASH", "gang"), ("SUP", "sup")):
        for s in (1, 2, 3):
            key = f"{pre}_{s}"
            groups = {}
            for d in (1, 2):
                q = rep[src][f"{s}D{d}"]
                groups[f"{s}D{d}" in ph] = groups.get(f"{s}D{d}" in ph, D(0)) + q
            # remove qty moved to extra blocks from the non-PH bucket first
            left = ded.get(key, D(0))
            for flag in (False, True):
                take = min(left, groups.get(flag, D(0)))
                if take:
                    groups[flag] -= take
                    left -= take
            for flag in (False, True):
                if groups.get(flag):
                    main.append(svc_line(key, groups[flag], ph=flag))
    for key in ("MECH", "FL4_OP", "FL5", "FL7", "FL16", "FL_OP", "TM_OP", "TRAILER"):
        q = rep["equip"][key] - ded.get(key, D(0))
        if q > 0:
            main.append(svc_line(key, q))
    if rep["nt_tugmaster"] > 0:
        main.append(svc_line("TM_RENT", rep["nt_tugmaster"]))
    for key in ("LOCK", "TOWING"):
        q, amt = rep["misc"][key]
        if amt > 0:
            q = q or D(1)
            s = K["services"][key]
            main.append(line(key, s["desc"], q, "Unit" if q == 1 else "Units", money(amt / q), s["tariff"], amount=amt, edi_unit="UNIT"))
    litres, amt = rep["misc"]["FUEL"]
    if amt > 0:
        s = K["services"]["DIESEL"]
        main.append(line("DIESEL", f"Supply of Diesel {litres:.2f} Litres", 1, "Unit", amt, s["tariff"],
                         amount=amt, edi_unit="UNIT"))
    blocks = [{"title": None, "lines": main}]
    for b in job.get("extra_blocks", []):
        ls = []
        for it in b["items"]:
            ls.append(svc_line(it["key"], it["qty"], ph=it.get("ph", False), rate=it.get("rate"), desc=it.get("desc")))
        blocks.append({"title": b["title"], "lines": ls})
    return blocks


# ---------------------------------------------------------------- ECL line items
def ecl_blocks(job):
    E = RATES["ecl"]
    blocks = []
    for sec in ("DIS", "LOD", "SHF", "LNR"):
        ls = []
        for code in ("LT2", "2TO5", "GT5", "GC"):
            q = D(job["cargo"].get(sec, {}).get(code, 0))
            if q:
                c = E["cargo"][code]
                mult = 2 if sec == "LNR" else 1
                ls.append(line(f"{code}_{sec}", f"{fmt_qty(q)} {c['desc']}", q, "", D(c["rate"]) * mult, None))
        if ls:
            blocks.append({"title": E["sections"][sec], "lines": ls})
    ls = []
    for sh in job.get("supply", []):          # {"shift":2, "ph":false, "crew":{"SHIFTER":2,...}}
        rate = E["ph_rate"] if sh.get("ph") else E["shift_rate"][str(sh["shift"])]
        nth = {1: "1ST", 2: "2ND", 3: "3RD"}[sh["shift"]]
        for role in ("SHIFTER", "SIGNALLER", "RAMP", "WARDEN", "LABOURER"):
            q = D(sh["crew"].get(role, 0))
            if q:
                plural, singular = E["roles"][role]
                ls.append(line(f"{role}_{sh['shift']}",
                               f"{fmt_qty(q)} {plural} WORKING ON {nth} SHIFT @ ${rate:.2f} PER {singular} PER SHIFT"
                               + (" (SUN/PH)" if sh.get("ph") else ""), q, "", rate, None))
    if ls:
        blocks.append({"title": "SUPPLY", "lines": ls})
    for b in job.get("extra_blocks", []):
        blocks.append({"title": b["title"], "lines": [line("EXTRA", it["desc"], it["qty"], "", it["rate"], None) for it in b["items"]]})
    return blocks


# ---------------------------------------------------------------- words
ONES = "ZERO ONE TWO THREE FOUR FIVE SIX SEVEN EIGHT NINE TEN ELEVEN TWELVE THIRTEEN FOURTEEN FIFTEEN SIXTEEN SEVENTEEN EIGHTEEN NINETEEN".split()
TENS = "_ _ TWENTY THIRTY FORTY FIFTY SIXTY SEVENTY EIGHTY NINETY".split()


def words(n):
    n = int(n)
    if n < 20:
        return ONES[n]
    if n < 100:
        return TENS[n // 10] + ("" if n % 10 == 0 else " " + ONES[n % 10])
    if n < 1000:
        return ONES[n // 100] + " HUNDRED" + ("" if n % 100 == 0 else " " + words(n % 100))
    for div, name in ((10**9, "BILLION"), (10**6, "MILLION"), (1000, "THOUSAND")):
        if n >= div:
            rest = n % div
            return words(n // div) + " " + name + ("" if rest == 0 else " " + words(rest))


def amount_in_words(x):
    x = money(x)
    dollars, cents = int(x), int((x - int(x)) * 100)
    s = "DOLLARS : " + words(dollars)
    s += f" AND CENTS {words(cents)} ONLY" if cents else " ONLY"
    return s


# ---------------------------------------------------------------- PDF rendering
def render_invoice(job, blocks, out_pdf, draft_mark=None):
    client = job["client"]
    C = RATES[client]
    W, H = A4
    c = canvas.Canvas(out_pdf, pagesize=A4)
    c.setTitle(f"Invoice {job['invoice_no']} {job['vessel']} V.{job['voyage']}")
    lh = ImageReader(os.path.join(HERE, RATES["company"]["letterhead"]))
    iw, ih = lh.getSize()
    lw = 150 * mm
    c.drawImage(lh, (W - lw) / 2, H - 14 * mm - lw * ih / iw, lw, lw * ih / iw, mask="auto")
    if draft_mark:
        c.saveState(); c.setFont("Helvetica-Bold", 60); c.setFillGray(0.88)
        c.translate(W / 2, H / 2); c.rotate(35); c.drawCentredString(0, 0, draft_mark); c.restoreState()
    y = H - 14 * mm - lw * ih / iw - 12 * mm
    # right header block
    hdr = [("Invoice No", job["invoice_no"]), ("Date", job["invoice_date"]), ("Terms", C["terms"]),
           ("Ref No", C["ref_no"]), ("No. of pages", "1")]
    if client == "kline":
        hdr.append(("Vendor Code", C["vendor_code"]))
    c.setFont("Helvetica", 9.5)
    hy = y
    for k, val in hdr:
        c.drawString(128 * mm, hy, k); c.drawString(152 * mm, hy, ":"); c.drawString(156 * mm, hy, str(val)); hy -= 4.6 * mm
    # left address block
    ay = y - 12 * mm
    c.setFont("Helvetica-Bold", 9.5)
    if client == "kline":
        c.drawString(20 * mm, ay, f"ATTN :  {job['attn']}"); ay -= 4.6 * mm
        c.setFont("Helvetica", 9.5)
        for l in C["address"]:
            c.drawString(20 * mm, ay, l); ay -= 4.6 * mm
    else:
        for i, l in enumerate(C["address"]):
            c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 9.5)
            c.drawString(20 * mm, ay, l); ay -= 4.6 * mm
    y = min(ay, hy) - 6 * mm
    c.setFont("Helvetica-Bold", 10)
    if client == "kline":
        c.drawString(20 * mm, y, f"M.V.{job['vessel']} VOY. {job['voyage']}"); y -= 4.8 * mm
        c.setFont("Helvetica", 9.5)
        c.drawString(20 * mm, y, f"ETA: {job['eta']}"); y -= 4.6 * mm
        c.drawString(20 * mm, y, f"PSA Berth : {job['berth']}"); y -= 9 * mm
    else:
        c.drawString(20 * mm, y, f"RE:  M.V. {job['vessel']} VOY. {job['voyage']}"); y -= 4.8 * mm
        c.setFont("Helvetica", 9.5)
        c.drawString(20 * mm, y, f"ARRIVED :  {job['arrived']}"); y -= 4.6 * mm
        c.drawString(20 * mm, y, f"BERTH AT {job['berth']}"); y -= 9 * mm
    # table header
    xR = 190 * mm
    if client == "kline":
        cols = {"desc": 20 * mm, "qty": 128 * mm, "unit": 131 * mm, "rate": 165 * mm}
        c.setFont("Helvetica-Bold", 9.5)
        c.drawCentredString(62 * mm, y, "Description"); c.drawCentredString(130 * mm, y, "Shift / Gang")
        c.drawRightString(cols["rate"], y, "Unit Price"); c.drawRightString(xR, y, "Amount")
    else:
        c.setFont("Helvetica-Bold", 9.5)
        c.drawString(20 * mm, y, "Description"); c.drawRightString(xR, y, "Amount (S$)")
    y -= 1.8 * mm; c.setLineWidth(0.6); c.line(20 * mm, y, xR, y); y -= 5 * mm
    total = D(0)
    fs, lh_ = (9, 4.4 * mm)
    for b in blocks:
        if not b["lines"]:
            continue
        if b["title"]:
            y -= 2 * mm
            c.setFont("Helvetica-Bold", fs)
            for t in (b["title"] if isinstance(b["title"], list) else [b["title"]]):
                c.drawString(20 * mm, y, t); y -= lh_
        c.setFont("Helvetica", fs)
        for l in b["lines"]:
            total += l["amount"]
            if client == "kline":
                c.drawString(cols["desc"], y, l["desc"])
                c.drawRightString(cols["qty"], y, fmt_qty(l["qty"]))
                c.drawString(cols["unit"], y, l["unit"])
                c.drawRightString(cols["rate"], y, fmt_money(l["rate"]))
                c.drawRightString(xR, y, fmt_money(l["amount"]))
            else:
                c.drawString(20 * mm, y, l["desc"])
                c.drawRightString(xR, y, fmt_money(l["amount"], dollar=False))
            y -= lh_
    y -= 1 * mm; c.line(160 * mm, y + 3.2 * mm, xR, y + 3.2 * mm)
    c.setFont("Helvetica-Bold", 10)
    c.drawRightString(157 * mm, y - 1 * mm, "TOTAL :")
    c.drawRightString(xR, y - 1 * mm, fmt_money(total, dollar=(client == "kline")))
    c.line(160 * mm, y - 3 * mm, xR, y - 3 * mm); c.line(160 * mm, y - 3.8 * mm, xR, y - 3.8 * mm)
    y -= 12 * mm
    if client == "ecl":
        c.setFont("Helvetica", 9.5); c.drawString(20 * mm, y, amount_in_words(total)); y -= 10 * mm
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, max(y - 8 * mm, 40 * mm), RATES["company"]["signoff"])
    sy = max(y - 26 * mm, 22 * mm)
    c.setDash(1, 2); c.line(20 * mm, sy, 75 * mm, sy); c.setDash()
    c.setFont("Helvetica-Oblique", 9); c.drawString(32 * mm, sy - 4.5 * mm, RATES["company"]["dept"])
    c.showPage(); c.save()
    return money(total)


# ---------------------------------------------------------------- EDI
EDI_HEADER = ("PROCESS_DATE,BILL_TYPE,BILLING_COMPANY,BILL_NUMBER,BILL_ITEM_NUMBER,ACCOUNT_NUMBER,BILL_DATE,REF_NUMBER,"
              "CONTAINER_NUMBER,TARIFF_CODE,TARIFF_DESCRIPTION,RATE,UNIT_DESCRIPTION,BILLABLE_UNIT,AMOUNT,FULL_VESSEL_NAME,"
              "FULL_OUT_VOY_NUMBER,FULL_IN_VOY_NUMBER,ABBR_VESSEL_NAME,ABBR_OUT_VOY_NUMBER,ABBR_IN_VOY_NUMBER,LINE_CODE,"
              "GROSS_TONNAGE,LOA,SERVICE_ROUTE,IN_SERVICE_ROUTE,LAST_BTR_DATE,ATB_DATE,ATU_DATE,FIRST_ACTIVITY_DATE,"
              "LAST_ACTIVITY_DATE,CONNECTING_FULL_VSL_NAME,CONNECTING_FULL_OUT_VOY_NUMBER,CONNECTING_ABBR_VSL_NAME,"
              "CONNECTING_ABBR_VOY_NUMBER,CONNECTING_SERVICE_ROUTE,CONNECTING_IN_SERVICE_ROUTE,CONNECTING_VESSEL_COD_DATE,"
              "CONNECTING_VESSEL_ATB_DATE,SERVICE_START_DATE,SERVICE_END_DATE,LOCATION_FROM,LOCATION_TO,BERTH_NUMBER,"
              "SLOT_OPERATOR,LOAD_DISC_INDICATOR,FROM,TO,CNTR_TYPE,CNTR_SIZE,ISO_SIZE_TYPE,DG_IMO_CLASS,TRANSHIP_INDICATOR,"
              "DEPOT_INDICATOR,REASON_CODE,LADEN_STATUS,CNTR_OPERATOR,GST_INDICATOR,GST_PERCENTAGE,CURRENCY_CODE,"
              "EXCHANGE_RATE,ORG_CODE,CHARGE_CATEGORY,CHARGE_TYPE,CHARGE_CLASSIFICATION_1,CHARGE_CLASSIFICATION_2,"
              "CHARGE_DESCRIPTION,DESCRIPTION_LINE_1,DESCRIPTION_LINE_2,DISCOUNT_TARIFF_CODE,DISCOUNT_PERCENT,"
              "CUSTOMER_REF_1,CUSTOMER_REF_2,CUSTOMER_REF_3,CUSTOMER_REF_4,CUSTOMER_REF_5,CUSTOMER_REF_6").split(",")


def write_edi(job, blocks, out_csv):
    E = RATES["kline"]["edi"]
    proc = job["process_datetime"]           # dd/mm/yyyy HH:MM:SS
    n = 0
    rows = []
    for b in blocks:
        for l in b["lines"]:
            n += 1
            code, tdesc, cdesc = l["tariff"]
            r = {h: "" for h in EDI_HEADER}
            amt = money(l["edi_rate"] * l["edi_qty"]) if l["key"] not in ("DIESEL", "LOCK", "TOWING") else l["amount"]
            r.update({"PROCESS_DATE": proc, "BILL_TYPE": E["bill_type"], "BILLING_COMPANY": E["billing_company"],
                      "BILL_NUMBER": job["invoice_no"], "BILL_ITEM_NUMBER": n, "BILL_DATE": job["bill_date"],
                      "TARIFF_CODE": code, "TARIFF_DESCRIPTION": tdesc, "RATE": f"{money(l['edi_rate']):.2f}",
                      "UNIT_DESCRIPTION": l["edi_unit"], "BILLABLE_UNIT": f"{l['edi_qty']:.3f}" if l["edi_unit"] == "M3" else f"{l['edi_qty']:.2f}",
                      "AMOUNT": f"{amt:.2f}", "FULL_VESSEL_NAME": job["vessel"], "FULL_OUT_VOY_NUMBER": job["voyage"],
                      "FULL_IN_VOY_NUMBER": job["voyage"], "SERVICE_START_DATE": job["service_start"],
                      "SERVICE_END_DATE": job["service_end"], "GST_INDICATOR": E["gst_indicator"],
                      "GST_PERCENTAGE": E["gst_pct"], "CURRENCY_CODE": E["currency"], "EXCHANGE_RATE": E["fx"],
                      "CHARGE_DESCRIPTION": cdesc, "DESCRIPTION_LINE_1": cdesc})
            rows.append(r)
    with open(out_csv, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=EDI_HEADER, quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
        w.writeheader(); w.writerows(rows)
    return sum(D(r["AMOUNT"]) for r in rows), n


# ---------------------------------------------------------------- stack
def build_stack(parts, out_pdf):
    """parts: list of (pdf_path, [page indexes] or None)"""
    w = PdfWriter()
    for p, pages in parts:
        rd = PdfReader(p)
        for i in (pages if pages is not None else range(len(rd.pages))):
            w.add_page(rd.pages[i])
    with open(out_pdf, "wb") as f:
        w.write(f)
    return out_pdf


def sheet_to_pdf(xls_path, sheet, outdir):
    """Render one worksheet of a workbook to PDF via LibreOffice."""
    import openpyxl
    x = xls_to_xlsx(xls_path, outdir)
    wb = openpyxl.load_workbook(x)
    for n in list(wb.sheetnames):
        if n != sheet:
            del wb[n]
    ws = wb[sheet]
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    tmp = os.path.join(outdir, f"_{sheet}_only.xlsx")
    wb.save(tmp)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, tmp], check=True, capture_output=True)
    return tmp[:-5] + ".pdf"
